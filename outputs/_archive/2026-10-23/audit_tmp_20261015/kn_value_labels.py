"""Extract the KNHANES value-label codebook (变量解释.docx) into a tidy CSV,
then print the value labels for the variables we must not guess.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

import docx
import pandas as pd

DOCX = Path(
    r"E:\BaiduNetdiskDownload\所有数据库\11. KNHANES 韩国国民健康与营养调查\变量解释.docx"
)
OUT = Path(__file__).resolve().parent.parent / ".audit_tmp" / "kn_value_labels.csv"

d = docx.Document(str(DOCX))
rows = []
section = ""
for par in d.paragraphs:
    s = par.text.strip()
    if s and not s.startswith(("变量解释", "来自KNHANES")):
        # headings look like "2.12吸烟"
        if re.match(r"^\d+(\.\d+)*\s*\S", s) and len(s) < 40:
            section = s
# tables don't carry their section; use table order + a running heading map
heading_map = {}
cur = ""
for i, t in enumerate(d.tables):
    txt = t.rows[0].cells[0].text.strip() if t.rows else ""
    if "变量类型" not in " ".join(c.text for c in t.rows[0].cells):
        # header row is a section banner
        cur = txt
    heading_map[i] = cur

for i, t in enumerate(d.tables):
    hdr = [c.text.strip() for c in t.rows[0].cells]
    ncol = len(hdr)
    for r in t.rows[1:]:
        cells = [c.text.strip().replace("\n", " ") for c in r.cells]
        if len(cells) < 4:
            continue
        if len(cells) == 4:
            qno, vtype, vname, vexp = cells
            cat = ""
        else:
            qno, vtype, vname, vexp, cat = cells[:5]
        if vname in ("变量名称",) or not vname:
            continue
        rows.append({"table": i, "banner": heading_map.get(i, ""), "q_no": qno,
                     "v_type": vtype, "var": vname, "label": vexp, "value_label": cat})

cb = pd.DataFrame(rows)
OUT.parent.mkdir(parents=True, exist_ok=True)
cb.to_csv(OUT, index=False, encoding="utf-8-sig")
print(f"codebook rows: {len(cb)}  vars: {cb['var'].nunique()}  -> {OUT}")

# section index
print("\n=== section banners (tables) ===")
seen = []
for i, t in enumerate(d.tables):
    b = heading_map.get(i, "")
    if b and b not in seen:
        seen.append(b)
print(seen)


def show(var, maxn=30):
    sub = cb[cb["var"] == var]
    print(f"\n=== {var} ({len(sub)} rows) ===")
    if sub.empty:
        print("  not found")
        return
    for r in sub.head(maxn).itertuples():
        print(f"  {str(r.q_no):>6s} | {str(r.v_type):>8s} | {str(r.label)[:34]:34s} | {str(r.value_label)[:60]}")


for v in ["BS1_1", "BS3_1", "BS5", "BD1", "BD1_11", "BD2_1", "BD2", "BD2_31", "BD2_32",
          "BA2_2_1", "BA2_22", "age", "sex", "edu", "incm5",
          "E_DM", "E_DM_r", "E_AMD", "E_AMD_r", "E_GR", "E_Pr_1", "E_VS_DS",
          "E_DRP_R", "E_AMD_ELY_R", "E_Grv", "E_Gr_p", "DH2_lt", "DH2_dg",
          "HE_Uacid", "HE_hsCRP", "HE_vitD", "PM25"]:
    show(v)
