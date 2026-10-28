#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Build cover_letter.docx from cover_letter.md (the .md is the source of truth).

The previous version of this script owned the letter text and overwrote the .md
on every run; that clobbered a hand-edited, journal-ready letter on 2026-10-28.
It now reads the .md and never writes it.

Run: python scripts/build_cover_letter.py
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "outputs" / "submission" / "cover_letter.md"
OUT = ROOT / "outputs" / "submission" / "cover_letter.docx"


def _add_md_paragraph(doc: Document, text: str) -> None:
    """Render a markdown line, honouring **bold** runs."""
    p = doc.add_paragraph()
    for i, chunk in enumerate(re.split(r"(\*\*[^*]+\*\*)", text)):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            run = p.add_run(chunk[2:-2])
            run.bold = True
        else:
            p.add_run(chunk)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15


def main() -> int:
    if not SRC.exists():
        raise SystemExit(f"missing source: {SRC}")
    lines = SRC.read_text(encoding="utf-8").splitlines()

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    for raw in lines:
        line = raw.rstrip()
        if not line:
            continue
        if line.startswith("# "):
            h = doc.add_heading(line[2:].strip(), level=1)
            for r in h.runs:
                r.font.name = "Times New Roman"
            continue
        if line.startswith("**Journal:**") or line.startswith("**Article type:**") \
                or line.startswith("**Title:**") or line.startswith("**Authors:**"):
            _add_md_paragraph(doc, line)
            continue
        if line.startswith("> "):
            _add_md_paragraph(doc, line[2:].strip())
            continue
        _add_md_paragraph(doc, line)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    words = len(SRC.read_text(encoding="utf-8").split())
    print(f"cover letter docx written: {words} words -> {OUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
